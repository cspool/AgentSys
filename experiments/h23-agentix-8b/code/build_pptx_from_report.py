#!/usr/bin/env python3
"""Turn an h23 HTML report into a PPTX that keeps all of its text and figures.

The deck follows the document, not a template: the h2 sections become chapters,
and inside a chapter every paragraph, figure, table and reading-rule card becomes
slides in the order the document put them. Nothing is summarised away — long
prose is paged across slides rather than trimmed, because the report's argument
lives in the prose as much as in the figures.

Figures arrive two ways and both are handled: paper figures are base64 <img>
payloads, and the report's own charts are inline <svg>, rasterised here with
cairosvg because PowerPoint does not take SVG.
"""
from __future__ import annotations

import argparse
import base64
import io
import re
from pathlib import Path

import lxml.html
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.55)
BODY_W = W - 2 * MARGIN
INK = RGBColor(0x1A, 0x1D, 0x23)
MUTED = RGBColor(0x5A, 0x63, 0x72)
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
BG = RGBColor(0xFF, 0xFF, 0xFF)
BAND = RGBColor(0xF2, 0xF5, 0xF9)
CJK = "Microsoft YaHei"


def set_font(run, size, bold=False, color=INK, name=CJK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag.split(':')[1]}")
        if el is None:
            el = etree.SubElement(
                rpr, f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag.split(':')[1]}")
        el.set("typeface", name)


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    return s


def add_text(slide, text, left, top, width, height, size, bold=False,
             color=INK, align=PP_ALIGN.LEFT, spacing=1.25):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        set_font(p.add_run(), size, bold, color)
        p.runs[0].text = line
    return box


def header(slide, chapter, title=None):
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(0.09))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    add_text(slide, chapter, MARGIN, Inches(0.22), BODY_W, Inches(0.4), 13, False, MUTED)
    if title:
        add_text(slide, title, MARGIN, Inches(0.6), BODY_W, Inches(0.8), 24, True, INK)


# ---------------------------------------------------------------- extraction
CJK_SVG_FONT = "Noto Sans CJK SC, Noto Serif CJK SC, Droid Sans Fallback, sans-serif"
CAMEL_ATTRS = ("viewBox", "preserveAspectRatio", "gradientUnits", "gradientTransform",
               "patternUnits", "patternTransform", "markerWidth", "markerHeight",
               "refX", "refY", "spreadMethod", "textLength", "lengthAdjust",
               "clipPathUnits", "maskUnits", "startOffset", "baseFrequency")


def svg_to_png(el) -> bytes | None:
    """Rasterise an inline SVG.

    Two repairs are needed first. The HTML parser lowercases attribute names, so
    viewBox arrives as viewbox and the renderer sees an SVG with no dimensions at
    all — it fails silently and the figure vanishes from the deck. And the report
    sizes its charts with CSS (width:100%), which means nothing outside a browser,
    so explicit pixel dimensions are taken from the viewBox.
    """
    import cairosvg
    raw = etree.tostring(el, encoding="unicode")
    for camel in CAMEL_ATTRS:
        raw = re.sub(rf"(?<=[\s<]){camel.lower()}=", f"{camel}=", raw)
    if "xmlns" not in raw[:400]:
        raw = raw.replace("<svg", '<svg xmlns="http://www.w3.org/2000/svg"', 1)
    # The report's charts name no font, so the rasteriser falls back to a Latin
    # face and every Chinese label comes out as a tofu box. Name a fallback that
    # actually has the glyphs on this machine.
    if "font-family" not in raw[:400]:
        raw = raw.replace("<svg", f'<svg font-family="{CJK_SVG_FONT}"', 1)
    # The root style carries width:100%;height:auto for the browser. Outside one,
    # height:auto resolves to nothing, the viewport collapses and the chart
    # rasterises as an empty image of the right size. Drop the CSS sizing and let
    # the viewBox decide.
    m = re.match(r'<svg([^>]*)>', raw)
    if m and "style=" in m.group(1):
        head = m.group(0)
        fixed = re.sub(r'(?:width|height)\s*:\s*[^;"]+;?', '', head)
        raw = fixed + raw[len(head):]
    box = re.search(r'viewBox="([\d.\-\s]+)"', raw)
    if box and " width=" not in raw[:400]:
        parts = box.group(1).split()
        if len(parts) == 4:
            vw, vh = float(parts[2]), float(parts[3])
            raw = raw.replace("<svg", f'<svg width="{vw}" height="{vh}"', 1)
    try:
        return cairosvg.svg2png(bytestring=raw.encode(), scale=2.0,
                                background_color="white")
    except Exception as exc:
        print(f"[pptx] svg rasterise failed: {type(exc).__name__}: {exc}", flush=True)
        return None


def img_to_png(el) -> bytes | None:
    src = el.get("src", "")
    if not src.startswith("data:"):
        return None
    try:
        return base64.b64decode(src.split(",", 1)[1])
    except Exception:
        return None


def extract_blocks(root):
    """Walk the tree in document order, classifying as we go.

    Recursion stops at a figure, table, svg or img: those own their subtree, and
    descending into them would emit their inner text twice. Identity-based
    de-duplication is not an option here — lxml creates element proxies on
    demand, so `id(el)` is not stable across a traversal.
    """
    blocks = []

    def push(kind, **kw):
        blocks.append({"kind": kind, **kw})

    def handle_figure(el):
        cap = clean(" ".join(c.text_content() for c in el.findall(".//figcaption")))
        png = None
        svg = el.find(".//svg")
        if svg is not None:
            png = svg_to_png(svg)
        if png is None:
            for im in el.findall(".//img"):
                png = img_to_png(im)
                if png:
                    break
        if png:
            push("figure", png=png, caption=cap)
        elif cap:
            push("para", text=cap)

    def handle_table(el):
        rows = []
        for tr in el.findall(".//tr"):
            cells = [clean(td.text_content())
                     for td in tr.findall("./th") + tr.findall("./td")]
            if cells:
                rows.append(cells)
        caps = el.findall(".//caption")
        if rows:
            push("table", rows=rows,
                 caption=clean(caps[0].text_content()) if caps else "")

    def walk(el):
        tag = el.tag if isinstance(el.tag, str) else ""
        if tag in ("script", "style", "head"):
            return
        if tag == "figure":
            handle_figure(el)
            return
        if tag == "table":
            handle_table(el)
            return
        if tag == "svg":
            png = svg_to_png(el)
            if png:
                push("figure", png=png, caption="")
            return
        if tag == "img":
            png = img_to_png(el)
            if png:
                push("figure", png=png, caption=clean(el.get("alt", "")))
            return
        if tag in ("h1", "h2", "h3"):
            txt = clean(el.text_content())
            if txt:
                push(tag, text=txt)
            return
        if tag in ("p", "li"):
            txt = clean(el.text_content())
            if txt and len(txt) > 1:
                push("para", text=txt)
            return
        for child in el:
            walk(child)

    walk(root)
    return blocks


# ---------------------------------------------------------------- rendering
def page_text(text, per_slide):
    """Split prose across slides on sentence boundaries, never mid-sentence."""
    parts = re.split(r"(?<=[。；！？.;!?])\s*", text)
    out, cur, n = [], [], 0
    for s in parts:
        if not s:
            continue
        if n + len(s) > per_slide and cur:
            out.append("".join(cur))
            cur, n = [], 0
        cur.append(s)
        n += len(s)
    if cur:
        out.append("".join(cur))
    return out or [text]


def add_figure_slide(prs, chapter, png, caption, index):
    """Place a figure, slicing a very tall one across slides.

    Some report charts are stacked timelines several times taller than they are
    wide. Fitting one of those into a 16:9 slide shrinks it until the labels are
    unreadable, which defeats the point of carrying the figure at all, so a tall
    figure is cut into slide-shaped bands and each band gets its own slide.
    """
    im = Image.open(io.BytesIO(png))
    iw, ih = im.size
    cap_h = Inches(0.95) if caption else Inches(0.05)
    top = Inches(0.9)
    avail_h = H - top - cap_h - Inches(0.25)
    avail_w = BODY_W
    max_ratio = avail_h / avail_w
    bands = max(1, min(6, int(round((ih / iw) / max_ratio))))
    band_px = ih // bands
    for b in range(bands):
        slide = blank(prs)
        header(slide, chapter)
        crop = im.crop((0, b * band_px, iw,
                        ih if b == bands - 1 else (b + 1) * band_px))
        cw, ch = crop.size
        scale = min(avail_w / cw, avail_h / ch)
        w, h = int(cw * scale), int(ch * scale)
        tmp = Path(f"/tmp/claude-0/pptx_fig_{index}_{b}.png")
        tmp.parent.mkdir(parents=True, exist_ok=True)
        crop.save(tmp)
        slide.shapes.add_picture(str(tmp), int((W - w) / 2), top,
                                 width=Emu(w), height=Emu(h))
        tmp.unlink(missing_ok=True)
        if caption:
            text = caption if bands == 1 else f"{caption}（{b + 1}/{bands}）"
            add_text(slide, text, MARGIN, top + Emu(h) + Inches(0.15),
                     BODY_W, cap_h, 13, False, MUTED)


def add_table_slide(prs, chapter, rows, caption):
    head, body = rows[0], rows[1:]
    cols = max(len(r) for r in rows)
    per = 11
    for page in range(0, max(len(body), 1), per):
        chunk = body[page:page + per]
        slide = blank(prs)
        header(slide, chapter)
        title = caption or "表"
        if len(body) > per:
            title = f"{title}（{page // per + 1}/{(len(body) - 1) // per + 1}）"
        add_text(slide, title, MARGIN, Inches(0.62), BODY_W, Inches(0.5), 17, True, INK)
        n_rows = len(chunk) + 1
        top = Inches(1.25)
        height = min(H - top - Inches(0.4), Inches(0.34) * n_rows)
        shape = slide.shapes.add_table(n_rows, cols, MARGIN, top, BODY_W, height)
        tbl = shape.table
        for c in range(cols):
            cell = tbl.cell(0, c)
            cell.text = head[c] if c < len(head) else ""
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    set_font(r, 11, True, INK)
        for r_i, row in enumerate(chunk, start=1):
            for c in range(cols):
                cell = tbl.cell(r_i, c)
                cell.text = row[c] if c < len(row) else ""
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        set_font(run, 10.5, False, INK)


def build(blocks, title, subtitle, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    cover = blank(prs)
    band = cover.shapes.add_shape(1, 0, Inches(2.3), W, Inches(2.4))
    band.fill.solid()
    band.fill.fore_color.rgb = BAND
    band.line.fill.background()
    add_text(cover, title, MARGIN, Inches(2.65), BODY_W, Inches(1.1), 34, True, INK)
    add_text(cover, subtitle, MARGIN, Inches(3.75), BODY_W, Inches(0.8), 15, False, MUTED)

    chapter = title
    fig_i = 0
    pending: list[str] = []

    def flush():
        nonlocal pending
        if not pending:
            return
        text = "\n".join(pending)
        for page in page_text(text, 700):
            slide = blank(prs)
            header(slide, chapter)
            add_text(slide, page, MARGIN, Inches(0.95), BODY_W,
                     H - Inches(1.4), 15, False, INK, spacing=1.45)
        pending = []

    for b in blocks:
        k = b["kind"]
        if k in ("h1", "h2"):
            flush()
            chapter = b["text"]
            if k == "h2":
                slide = blank(prs)
                bar = slide.shapes.add_shape(1, 0, Inches(3.15), W, Inches(1.35))
                bar.fill.solid()
                bar.fill.fore_color.rgb = BAND
                bar.line.fill.background()
                add_text(slide, b["text"], MARGIN, Inches(3.4), BODY_W, Inches(1.0),
                         26, True, INK)
        elif k == "h3":
            flush()
            slide = blank(prs)
            header(slide, chapter, b["text"])
        elif k == "para":
            pending.append(b["text"])
        elif k == "figure":
            flush()
            fig_i += 1
            add_figure_slide(prs, chapter, b["png"], b["caption"], fig_i)
        elif k == "table":
            flush()
            add_table_slide(prs, chapter, b["rows"], b.get("caption", ""))
    flush()
    prs.save(out_path)
    return len(prs.slides._sldIdLst), fig_i


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--html", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--title", required=True)
    ap.add_argument("--subtitle", default="")
    a = ap.parse_args()
    doc = lxml.html.fromstring(a.html.read_bytes())
    body = doc.find("body") if doc.find("body") is not None else doc
    blocks = extract_blocks(body)
    kinds = {}
    for b in blocks:
        kinds[b["kind"]] = kinds.get(b["kind"], 0) + 1
    slides, figs = build(blocks, a.title, a.subtitle, str(a.out))
    print({"blocks": kinds, "slides": slides, "figures_placed": figs,
           "out": str(a.out), "size_kb": round(a.out.stat().st_size / 1024)})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
